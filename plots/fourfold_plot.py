"""
Presentation av fyrfältstabell: färgad HTML för skärm och Excel för nedladdning.

Färgsättning
------------
Diagonalen (överensstämmelse) grön, off-diagonalen (avvikelser) röd,
med mörkare ton ju fler observationer. Detta gör att ögat direkt hittar
felklassificeringarna, vilket är hela poängen med tabellen.
"""

import io
from typing import Dict

# Färger
GREEN_BG = "D1FAE5"
GREEN_STRONG = "6EE7B7"
RED_BG = "FEE2E2"
RED_STRONG = "FCA5A5"
HDR_BG = "1E40AF"
LBL_BG = "DCE6F1"
GREY_BG = "F1F5F9"


def _shade(count: int, total: int, weak: str, strong: str) -> str:
    """Ljusare färg vid få observationer, mörkare vid många."""
    if total <= 0 or count == 0:
        return "FFFFFF"
    frac = count / total
    return strong if frac >= 0.15 else weak


def _f(v, d=1):
    if v is None:
        return ""
    try:
        if v != v:          # NaN
            return "–"
    except Exception:
        pass
    return f"{v:.{d}f}".replace(".", ",")


def _ci(t, d=1):
    lo, hi = t
    return f"{_f(lo,d)}–{_f(hi,d)}"


def fourfold_html(res: Dict, ref_name="Reference method", cand_name="Candidate method",
                  pos_label="Positive", neg_label="Negative", t=None) -> str:
    if t is None:
        t = lambda x: x
    """Färgad 2×2-tabell som HTML, för visning i appen."""
    n = res["n"]
    cells = {
        "tp": (res["tp"], _shade(res["tp"], n, GREEN_BG, GREEN_STRONG)),
        "fp": (res["fp"], _shade(res["fp"], n, RED_BG, RED_STRONG)),
        "fn": (res["fn"], _shade(res["fn"], n, RED_BG, RED_STRONG)),
        "tn": (res["tn"], _shade(res["tn"], n, GREEN_BG, GREEN_STRONG)),
    }
    td = ("padding:14px 10px;text-align:center;font-size:1.35rem;"
          "font-weight:700;color:#0F172A;border:1px solid #CBD5E1;")
    lbl = ("padding:10px;text-align:center;font-weight:600;font-size:0.8rem;"
           "background:#DCE6F1;border:1px solid #CBD5E1;color:#1E293B;")
    tot = ("padding:10px;text-align:center;font-weight:600;font-size:0.95rem;"
           "background:#F1F5F9;border:1px solid #CBD5E1;color:#334155;")
    sub = "display:block;font-size:0.62rem;font-weight:500;color:#475569;margin-top:3px;"

    return f"""
<table style="border-collapse:collapse;margin:6px 0 2px 0;">
  <tr>
    <td style="border:none"></td>
    <td colspan="3" style="{lbl}background:#1E40AF;color:#fff;">{ref_name}</td>
  </tr>
  <tr>
    <td style="border:none"></td>
    <td style="{lbl}">{pos_label}</td>
    <td style="{lbl}">{neg_label}</td>
    <td style="{tot}">{t("Total")}</td>
  </tr>
  <tr>
    <td style="{lbl}writing-mode:vertical-rl;transform:rotate(180deg);
               padding:6px;background:#1E40AF;color:#fff;" rowspan="2">{cand_name}</td>
    <td style="{td}background:#{cells['tp'][1]};">{cells['tp'][0]}
        <span style="{sub}">{t("True positive")}</span></td>
    <td style="{td}background:#{cells['fp'][1]};">{cells['fp'][0]}
        <span style="{sub}">{t("False positive")}</span></td>
    <td style="{tot}">{res['n_cand_pos']}</td>
  </tr>
  <tr>
    <td style="{td}background:#{cells['fn'][1]};">{cells['fn'][0]}
        <span style="{sub}">{t("False negative")}</span></td>
    <td style="{td}background:#{cells['tn'][1]};">{cells['tn'][0]}
        <span style="{sub}">{t("True negative")}</span></td>
    <td style="{tot}">{res['n_cand_neg']}</td>
  </tr>
  <tr>
    <td style="border:none"></td>
    <td style="{tot}">{res['n_ref_pos']}</td>
    <td style="{tot}">{res['n_ref_neg']}</td>
    <td style="{tot}font-weight:700;">{n}</td>
  </tr>
</table>
"""


def build_fourfold_excel(res: Dict, ref_name="Reference method",
                         cand_name="Candidate method", analyte="",
                         pos_label="Positive", neg_label="Negative",
                         warnings=None, t=None) -> bytes:
    if t is None:
        t = lambda x: x
    """Excel-arbetsbok: färgad fyrfältstabell + måtttabell + tolkningshjälp."""
    import openpyxl
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

    thin = Side(style="thin", color="94A3B8")
    med = Side(style="medium", color="1E293B")
    B = Border(left=thin, right=thin, top=thin, bottom=thin)
    CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
    RIGHT = Alignment(horizontal="right", vertical="center")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Fyrfältstabell"
    n = res["n"]
    ref_mode = res["mode"] == "reference"

    def style_only(r, c, fill=None, border=True):
        """Styla en cell utan att skriva varde - kravs for merged cells."""
        cell = ws.cell(row=r, column=c)
        if fill:
            cell.fill = PatternFill("solid", fgColor=fill)
        if border:
            cell.border = B
        return cell

    def put(r, c, v, fill=None, bold=False, size=11, border=True,
            align=CENTER, color="000000"):
        cell = ws.cell(row=r, column=c, value=v)
        if fill:
            cell.fill = PatternFill("solid", fgColor=fill)
        cell.font = Font(bold=bold, size=size, color=color)
        cell.alignment = align
        if border:
            cell.border = B
        return cell

    # ── Rubrik ───────────────────────────────────────────────────────────────
    ws.merge_cells("A1:F1")
    put(1, 1, f"{t('Fourfold table')}{(' – ' + analyte) if analyte else ''}",
        fill=HDR_BG, bold=True, size=14, color="FFFFFF")
    ws.row_dimensions[1].height = 26

    # ── 2×2-tabellen ─────────────────────────────────────────────────────────
    # Kolumnrubrikerna MASTE ligga over sina datakolumner:
    #   D = positiv, E = negativ, F = summa
    ws.merge_cells("D3:E3")
    put(3, 4, ref_name, fill=HDR_BG, bold=True, color="FFFFFF")
    put(3, 3, "", fill=None, border=False)
    put(3, 6, "", fill=None, border=False)
    put(4, 3, "", fill=None, border=False)
    put(4, 4, pos_label, fill=LBL_BG, bold=True)
    put(4, 5, neg_label, fill=LBL_BG, bold=True)
    put(4, 6, t("Total"), fill=GREY_BG, bold=True)

    ws.merge_cells("B5:B6")
    put(5, 2, cand_name, fill=HDR_BG, bold=True, color="FFFFFF")
    put(5, 1, "", border=False)

    put(5, 3, pos_label, fill=LBL_BG, bold=True)
    put(6, 3, neg_label, fill=LBL_BG, bold=True)

    # Placera cellerna med färg
    tp = put(5, 4, res["tp"], fill=_shade(res["tp"], n, GREEN_BG, GREEN_STRONG),
             bold=True, size=14)
    fp = put(5, 5, res["fp"], fill=_shade(res["fp"], n, RED_BG, RED_STRONG),
             bold=True, size=14)
    fn = put(6, 4, res["fn"], fill=_shade(res["fn"], n, RED_BG, RED_STRONG),
             bold=True, size=14)
    tn = put(6, 5, res["tn"], fill=_shade(res["tn"], n, GREEN_BG, GREEN_STRONG),
             bold=True, size=14)
    for c in (tp, fp, fn, tn):
        c.border = Border(left=med, right=med, top=med, bottom=med)

    put(7, 4, res["n_ref_pos"], fill=GREY_BG, bold=True)
    put(7, 5, res["n_ref_neg"], fill=GREY_BG, bold=True)
    put(5, 6, res["n_cand_pos"], fill=GREY_BG, bold=True)
    put(6, 6, res["n_cand_neg"], fill=GREY_BG, bold=True)
    put(7, 6, n, fill=GREY_BG, bold=True, size=12)
    put(7, 3, t("Total"), fill=GREY_BG, bold=True)

    # Cellförklaring
    put(9, 4, t("Green = agreement"), fill=GREEN_BG, size=9)
    put(9, 5, t("Red = disagreement"), fill=RED_BG, size=9)

    # ── Måtttabell ───────────────────────────────────────────────────────────
    r0 = 11
    ws.merge_cells(f"B{r0}:F{r0}")
    put(r0, 2, t("Results"), fill=HDR_BG, bold=True, size=12, color="FFFFFF")
    ws.merge_cells(f"B{r0+1}:C{r0+1}")
    put(r0 + 1, 2, t("Measure"), fill=LBL_BG, bold=True, size=10)
    style_only(r0 + 1, 3, fill=LBL_BG)
    put(r0 + 1, 4, t("Value"), fill=LBL_BG, bold=True, size=10)
    put(r0 + 1, 5, f"{int(res['conf']*100)} % {t('CI')}", fill=LBL_BG, bold=True, size=10)
    put(r0 + 1, 6, t("Count"), fill=LBL_BG, bold=True, size=10)

    if ref_mode:
        rows = [
            (t("Sensitivity"),       f"{_f(res['ppa'])} %", _ci(res["ppa_ci"]),
             f"{res['tp']}/{res['n_ref_pos']}"),
            (t("Specificity"),       f"{_f(res['npa'])} %", _ci(res["npa_ci"]),
             f"{res['tn']}/{res['n_ref_neg']}"),
            (t("Positive predictive value"), f"{_f(res['ppv'])} %", _ci(res["ppv_ci"]),
             f"{res['tp']}/{res['n_cand_pos']}"),
            (t("Negative predictive value"), f"{_f(res['npv'])} %", _ci(res["npv_ci"]),
             f"{res['tn']}/{res['n_cand_neg']}"),
            (t("Correctly classified"), f"{_f(res['opa'])} %", _ci(res["opa_ci"]),
             f"{res['tp']+res['tn']}/{n}"),
            (t("Prevalence in the material"), f"{_f(res['prevalence'])} %",
             _ci(res["prevalence_ci"]), f"{res['n_ref_pos']}/{n}"),
            (t("Positive likelihood ratio (LR+)"), _f(res["lr_pos"], 2), "", ""),
            (t("Negative likelihood ratio (LR−)"), _f(res["lr_neg"], 2), "", ""),
        ]
    else:
        rows = [
            (t("Positive percent agreement (PPA)"),
             f"{_f(res['ppa'])} %", _ci(res["ppa_ci"]),
             f"{res['tp']}/{res['n_ref_pos']}"),
            (t("Negative percent agreement (NPA)"),
             f"{_f(res['npa'])} %", _ci(res["npa_ci"]),
             f"{res['tn']}/{res['n_ref_neg']}"),
            (t("Overall percent agreement (OPA)"),
             f"{_f(res['opa'])} %", _ci(res["opa_ci"]),
             f"{res['tp']+res['tn']}/{n}"),
        ]

    rows += [
        (t("Cohen's kappa"), _f(res["kappa"], 3), _ci(res["kappa_ci"], 3),
         res["kappa_tolkning"]),
        (t("McNemar test (p-value)"),
         ("< 0,001" if res["mcnemar_p"] < 0.001 else _f(res["mcnemar_p"], 3)),
         res["mcnemar_metod"],
         t("systematic difference") if res["mcnemar_p"] < 0.05 else t("none detected")),
    ]

    for i, row in enumerate(rows):
        rr = r0 + 2 + i
        is_key = i < 2
        fill = "FFF9DB" if is_key else ("FFFFFF" if i % 2 else "F8FAFC")
        ws.merge_cells(f"B{rr}:C{rr}")
        put(rr, 2, row[0], fill=fill, bold=is_key, size=10, align=RIGHT)
        style_only(rr, 3, fill=fill)
        put(rr, 4, row[1], fill=fill, bold=is_key, size=10)
        put(rr, 5, row[2], fill=fill, size=9)
        put(rr, 6, row[3], fill=fill, size=9)

    # ── Varningar ────────────────────────────────────────────────────────────
    rw = r0 + 3 + len(rows) + 1
    if warnings:
        ws.merge_cells(f"B{rw}:F{rw}")
        put(rw, 2, t("Points to consider"), fill="FEF3C7", bold=True, size=11)
        for k, txt in enumerate(warnings):
            ws.merge_cells(f"B{rw+1+k}:F{rw+1+k}")
            put(rw + 1 + k, 2, "• " + txt, fill="FFFBEB", size=9,
                align=Alignment(horizontal="left", vertical="center",
                                wrap_text=True))
            ws.row_dimensions[rw + 1 + k].height = 26
        rw += len(warnings) + 2

    # ── Fotnot ───────────────────────────────────────────────────────────────
    notes = []
    if ref_mode:
        notes.append("Sensitivitet och specificitet förutsätter att "
                     "referensmetoden utgör facit för sant tillstånd.")
        notes.append("Prediktiva värden gäller endast vid prevalensen i detta "
                     "material och kan inte överföras till en population med "
                     "annan prevalens.")
    else:
        notes.append("Ingen av metoderna antas utgöra facit. Därför redovisas "
                     "procentuell överensstämmelse och inte sensitivitet och "
                     "specificitet (CLSI EP12-A2, FDA 2007).")
    notes += [
        "Konfidensintervall beräknade med Wilsons score-metod.",
        "McNemars test prövar om metoderna skiljer sig systematiskt åt; "
        "endast diskordanta par bidrar.",
        "Cohens kappa beskriver överensstämmelse korrigerad för slumpen.",
    ]
    try:
        import sys as _s, os as _o
        _s.path.insert(0, _o.path.dirname(_o.path.dirname(_o.path.abspath(__file__))))
        from version import stamp as _st
        notes.append(_st())
    except Exception:
        pass

    for k, t in enumerate(notes):
        ws.merge_cells(f"B{rw+k}:F{rw+k}")
        c = ws.cell(row=rw + k, column=2, value=t)
        c.font = Font(size=8, italic=True, color="64748B")
        c.alignment = Alignment(horizontal="left", wrap_text=True)
        ws.row_dimensions[rw + k].height = 22

    # B rymmer de langa mattnamnen, C-F tabellcellerna
    for col, w in zip("ABCDEF", [3, 23, 14, 14, 17, 13]):
        ws.column_dimensions[col].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()

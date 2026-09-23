"""
Översikt: jämför ALLA kopplade analyser mellan två filer i en körning.

Ingen ny statistik införs. Varje rad beräknas med samma validerade
funktioner som den enskilda analysen (match_two_files, passing_bablok,
deming, weighted_deming, summary_stats), så resultatet för en analys är
identiskt med det man får genom att välja den enskilt. Det verifieras i
tests/instrument/testa_oversikt.py.
"""
import io
import re
from typing import Callable, Dict, List, Optional, Tuple

import numpy as np
import pandas as pd

from analysis.data_loader import match_two_files, list_analytes
from analysis.file_reader import suggest_analyte
from analysis.regression import passing_bablok
from analysis.deming import deming, weighted_deming
from analysis.statistics import summary_stats

MIN_RECOMMENDED = 40          # CLSI EP09: minst 40 prov rekommenderas


def suggested_pairs(df_a, an_col_a, df_b, an_col_b) -> List[Tuple[str, str, str]]:
    """(analys i A, föreslagen analys i B, grund) för varje analys i A med ett förslag."""
    lb = list_analytes(df_b, an_col_b)
    out = []
    for a in list_analytes(df_a, an_col_a):
        b, why = suggest_analyte(a, lb)
        if b is not None:
            out.append((a, b, why))
    return out


def _fit(x, y, method, error_ratio, weighted):
    if method == "Deming":
        return (weighted_deming(x, y, error_ratio=error_ratio) if weighted
                else deming(x, y, error_ratio=error_ratio))
    return passing_bablok(x, y)


def compare_all(df_a, df_b, id_a, id_b, an_col_a, an_col_b, res_a, res_b, pairs,
                method: str = "Passing–Bablok", error_ratio: float = 1.0,
                weighted: bool = False, ignore_leading_zeros: bool = True,
                dup_strategy: str = "first_valid", label_a: str = "A", label_b: str = "B",
                progress: Optional[Callable[[int, int, str], None]] = None
                ) -> Tuple[pd.DataFrame, Dict[str, pd.DataFrame]]:
    """
    Returnerar (översiktstabell, {analys: matchningsrapport}).
    Kolumnnamnen är engelska nycklar som översätts vid visning.
    """
    rows, details = [], {}
    for i, (a, b, why) in enumerate(pairs):
        if progress:
            progress(i, len(pairs), a)
        label = a if a == b else f"{a} ↔ {b}"
        row = {"Analysis A": a, "Analysis B": b, "Pairing": why, "Pairs used": 0,
               "Not used": 0, "Slope": np.nan, "Slope 95% CI": "", "Intercept": np.nan,
               "Intercept 95% CI": "", "Mean bias": np.nan, "Mean bias (%)": np.nan,
               "r": np.nan, "Slope CI includes 1": "", "Intercept CI includes 0": "",
               "Note": ""}
        try:
            x, y, rep, sm = match_two_files(
                df_a, df_b, id_a, id_b, an_col_a, an_col_b, res_a, res_b, a, b,
                label_a, label_b, ignore_leading_zeros=ignore_leading_zeros,
                dup_strategy=dup_strategy)
            details[label] = rep
            row["Pairs used"], row["Not used"] = sm["used"], sm["excluded"]
            notes = []
            if sm.get("unit_factor"):
                notes.append(f"unit difference ×{sm['unit_factor']:g}?")
            if len(x) < 3:
                notes.append("too few pairs")
            else:
                r = _fit(x, y, method, error_ratio, weighted)
                st = summary_stats(x, y)
                row.update({
                    "Slope": r["slope"],
                    "Slope 95% CI": (r["slope_lower"], r["slope_upper"]),
                    "Intercept": r["intercept"],
                    "Intercept 95% CI": (r["intercept_lower"], r["intercept_upper"]),
                    "Mean bias": st["mean_diff"],
                    "Mean bias (%)": float(np.mean((y - x) / ((x + y) / 2)) * 100)
                    if np.all((x + y) != 0) else np.nan,
                    "r": st["pearson_r"],
                    "Slope CI includes 1": "yes" if r["slope_lower"] <= 1 <= r["slope_upper"] else "no",
                    "Intercept CI includes 0": "yes" if r["intercept_lower"] <= 0 <= r["intercept_upper"] else "no",
                })
                if len(x) < MIN_RECOMMENDED:
                    notes.append(f"fewer than {MIN_RECOMMENDED} pairs")
            row["Note"] = "; ".join(notes)
        except Exception as e:                              # en analys stoppar inte resten
            row["Note"] = f"error: {e}"
        rows.append(row)
    return pd.DataFrame(rows), details


def _sheet_name(name: str, used: set) -> str:
    base = re.sub(r"[\[\]\*\?/\\:]", "-", str(name))[:28] or "Analys"
    out, k = base, 2
    while out.lower() in used:
        out = f"{base[:25]}~{k}"; k += 1
    used.add(out.lower())
    return out


def build_overview_excel(overview: pd.DataFrame, details: Dict[str, pd.DataFrame],
                         t: Callable[[str], str] = lambda s: s, decimals: int = 3,
                         stamp: str = "") -> bytes:
    """Excel: översiktsflik + en flik per analys med alla par och orsaker."""
    from openpyxl.styles import Font, PatternFill, Alignment
    fmt = lambda v: "" if v is None or (isinstance(v, float) and np.isnan(v)) else \
        f"{v:.{decimals}f}".replace(".", ",") if isinstance(v, (float, np.floating)) else v
    ci = lambda c: "" if not c else f"{fmt(float(c[0]))} – {fmt(float(c[1]))}"
    ov = overview.copy()
    for c in ("Slope 95% CI", "Intercept 95% CI"):
        ov[c] = ov[c].map(ci)
    for c in ("Slope", "Intercept", "Mean bias", "Mean bias (%)", "r"):
        ov[c] = ov[c].map(fmt)
    for c in ("Pairing", "Slope CI includes 1", "Intercept CI includes 0"):
        ov[c] = ov[c].map(lambda v: t(v) if v else v)
    ov["Note"] = ov["Note"].map(lambda v: "; ".join(_t_note(p, t) for p in v.split("; ")) if v else v)
    ov.columns = [t(c) for c in ov.columns]

    buf = io.BytesIO()
    head_fill = PatternFill("solid", fgColor="1E3A8A")
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        sh = t("Overview")[:31]
        ov.to_excel(w, sheet_name=sh, index=False, startrow=1)
        ws = w.sheets[sh]
        ws.cell(row=1, column=1, value=stamp).font = Font(italic=True, size=8, color="64748B")
        for c in ws[2]:
            c.font = Font(bold=True, color="FFFFFF"); c.fill = head_fill
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for i, wd in enumerate([16, 16, 12, 10, 10, 10, 18, 10, 18, 11, 12, 8, 12, 12, 30]):
            ws.column_dimensions[chr(65 + i)].width = wd
        ws.freeze_panes = "C3"
        used = {sh.lower()}
        for label, rep in details.items():
            name = _sheet_name(label, used)
            rep.to_excel(w, sheet_name=name, index=False)
            for c in w.sheets[name][1]:
                c.font = Font(bold=True)
    return buf.getvalue()


def _t_note(p: str, t) -> str:
    m = re.match(r"unit difference ×(.+)\?$", p)
    if m:
        return t("unit difference ×{f}?").format(f=m.group(1))
    m = re.match(r"fewer than (\d+) pairs$", p)
    if m:
        return t("fewer than {n} pairs").format(n=m.group(1))
    return t(p) if not p.startswith("error:") else p

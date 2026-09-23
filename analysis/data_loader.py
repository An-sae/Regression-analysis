"""
Data Loader — Long-format CSV matching and precision extraction
CLSI EP15-A3 compatible; matches SampleID x Analyte across two files.
"""
import io, numpy as np, pandas as pd
from collections import Counter
from typing import Dict, List, Optional, Tuple


from analysis.file_reader import (read_table, list_sheets, normalize_ids,
                                  parse_numeric, suggest_analyte, guess_column,
                                  canonical_analyte)


class MultipleResultsError(ValueError):
    """Flera resultat per prov-ID utan analyskolumn: parning vore gissning."""


def load_long_format(raw_bytes: bytes, fname: str, has_header: bool = True,
                     sheet=None, header_row="auto") -> pd.DataFrame:
    """
    Läs en instrument- eller LIS-export. Alla celler läses som TEXT.
    Kodning, avgränsare, rubrikrad och Excel-blad identifieras automatiskt;
    hur filen tolkades finns i df.attrs["read_info"].

    has_header=False -> ingen rubrikrad; kolumnerna heter "Column 1", ...
    """
    try:
        df, info = read_table(raw_bytes, fname, sheet=sheet,
                              header_row=header_row if has_header else None)
    except Exception as e:
        raise ValueError(f"Could not read file: {e}")
    df.attrs["read_info"] = info
    return df


def _coerce_numeric(series: pd.Series) -> pd.Series:
    """Bakåtkompatibel: tolka en kolumn till tal med kolumnvis decimaltecken."""
    vals, _ = parse_numeric(series)
    return vals["value"]


_SUMMARY_IDS = (r"(MEAN|MEDEL|MEDELV[ÄA]RDE|AVERAGE|AVG|MEDIAN|SD|STDAV|STD|STANDARDAVVIKELSE|"
                r"CV|CV%|VARIANS|VARIANCE|SUM|SUMMA|TOTAL|TOTALT|ANTAL|COUNT|N|MIN|MAX|"
                r"MINIMUM|MAXIMUM|RANGE)\.?:?")


def unit_factor(x: np.ndarray, y: np.ndarray) -> Optional[float]:
    """
    Trolig enhetsskillnad mellan filerna: om kvoten y/x ligger nära en tiopotens
    (±25 %), t.ex. g/L mot g/dL (10) eller fraktion mot procent (0,01),
    returneras tiopotensen. Annars None. Kräver minst fem positiva par.
    """
    x = np.asarray(x, float); y = np.asarray(y, float)
    ok = (x > 0) & (y > 0) & np.isfinite(x) & np.isfinite(y)
    if ok.sum() < 5:
        return None
    lr = float(np.log10(np.median(y[ok] / x[ok])))
    k = int(round(lr))
    if k != 0 and abs(lr - k) < np.log10(1.25):
        return float(10.0 ** k)
    return None


def _prepare(df, id_col, an_col, res_col, analyte, ignore_leading_zeros):
    """Filtrera analys, normalisera ID, tolka resultat. Returnerar ren tabell."""
    d = df
    if an_col not in (None, "N/A") and analyte not in (None, "ALL"):
        d = d[d[an_col].astype(str).str.strip() == str(analyte).strip()]
    d = d.copy()
    key, sci = normalize_ids(d[id_col], ignore_leading_zeros)
    parsed, pinfo = parse_numeric(d[res_col])
    out = pd.DataFrame({
        "key": key.values, "id_orig": d[id_col].astype(str).values,
        "raw": d[res_col].astype(str).values, "value": parsed["value"].values,
        "qualifier": parsed["qualifier"].values, "flag": parsed["flag"].values,
        "reason": parsed["reason"].values, "sci": sci.values,
    })
    # Summarader längst ned i LIS-rapporter ('Medelvärde', 'SD', 'Antal') är
    # inga prov och får aldrig paras ihop med varandra mellan filerna.
    summ = out["key"].str.fullmatch(_SUMMARY_IDS)
    info = {"parse": pinfo, "n_rows": len(out), "n_sci": int(sci.sum()),
            "n_blank_id": int(((key == "") & ~sci).sum()),
            "n_summary": int(summ.sum()),
            "summary_ids": sorted(set(out.loc[summ, "id_orig"].str.strip()))}
    out = out[(out["key"] != "") & ~summ].reset_index(drop=True)
    return out, info


def find_duplicates(df, id_col, analysis_col, result_col, ignore_leading_zeros=True) -> pd.DataFrame:
    """Rader där samma (normaliserade) prov-ID förekommer mer än en gång per analys."""
    key, _ = normalize_ids(df[id_col], ignore_leading_zeros)
    grp = [key] if analysis_col in (None, "N/A") else [key, df[analysis_col].astype(str)]
    counts = df.groupby(grp)[result_col].transform("count")
    dups = df[counts > 1].copy()
    if not dups.empty:
        dups["_occurrence"] = dups.groupby(grp[: len(grp)]).cumcount() + 1
    return dups


DUP_STRATEGIES = ("first_valid", "last_valid", "mean", "first", "last")


def _dedupe(t: pd.DataFrame, strategy: str) -> Tuple[pd.DataFrame, int]:
    """Välj ett resultat per nyckel. Returnerar (tabell, antal borttagna rader)."""
    n0 = len(t)
    if not t["key"].duplicated().any():
        return t, 0
    t = t.reset_index(drop=True)
    t["_ok"] = np.isfinite(t["value"].astype(float))
    if strategy == "mean":
        agg = t.groupby("key", sort=False).agg(
            id_orig=("id_orig", "first"), value=("value", "mean"),
            raw=("raw", lambda r: " | ".join(r)), qualifier=("qualifier", "first"),
            flag=("flag", "first"), sci=("sci", "first"),
            reason=("reason", lambda r: "" if (r == "").any() else r.iloc[0]))
        return agg.reset_index(), n0 - len(agg)
    if strategy in ("first_valid", "last_valid"):
        t = t.sort_values(["key", "_ok"], ascending=[True, False], kind="stable") \
            if strategy == "first_valid" else t
        if strategy == "last_valid":
            t = t.iloc[::-1].sort_values(["key", "_ok"], ascending=[True, False], kind="stable")
        out = t.drop_duplicates("key", keep="first")
    else:
        out = t.drop_duplicates("key", keep="last" if strategy == "last" else "first")
    return out.drop(columns="_ok"), n0 - len(out)


def get_common_analytes(df_a, df_b, analysis_col_a, analysis_col_b) -> List[str]:
    set_a = set(df_a[analysis_col_a].dropna().astype(str).str.strip().unique())
    set_b = set(df_b[analysis_col_b].dropna().astype(str).str.strip().unique())
    return sorted(set_a & set_b)


def list_analytes(df, analysis_col) -> List[str]:
    return sorted(v for v in df[analysis_col].dropna().astype(str).str.strip().unique() if v)


_REASON_TXT = {"below": "below measuring range", "above": "above measuring range",
               "text": "text result", "empty": "no result"}


def match_two_files(df_a, df_b, id_col_a, id_col_b, analysis_col_a, analysis_col_b,
                    result_col_a, result_col_b, analyte_a, analyte_b=None,
                    label_a="Method A", label_b="Method B",
                    ignore_leading_zeros=True, dup_strategy="first_valid",
                    single_analyte=False):
    """
    Matcha två filer på normaliserat prov-ID för analys analyte_a (fil A)
    och analyte_b (fil B; samma namn som A om None).

    Om analyskolumn saknas ("N/A") och något prov-ID förekommer flera gånger
    stoppas matchningen med MultipleResultsError, eftersom det inte går att
    avgöra vilka resultat som hör ihop — utom om användaren intygat att filen
    bara innehåller en analys (single_analyte=True); då behandlas upprepningar
    som omkörningar och löses med dup_strategy.

    Returnerar (x, y, rapport, sammanfattning).
    """
    if analyte_b is None:
        analyte_b = analyte_a
    A, ia = _prepare(df_a, id_col_a, analysis_col_a, result_col_a, analyte_a, ignore_leading_zeros)
    B, ib = _prepare(df_b, id_col_b, analysis_col_b, result_col_b, analyte_b, ignore_leading_zeros)

    for T_, col, nm in ((A, analysis_col_a, "A"), (B, analysis_col_b, "B")):
        if col in (None, "N/A") and not single_analyte and T_["key"].duplicated().any():
            n = int(T_["key"].duplicated(keep=False).sum())
            raise MultipleResultsError(
                f"File {nm}: {n} rows share a sample ID but no analysis column is "
                "selected, so results cannot be paired safely.")

    A, dup_a = _dedupe(A, dup_strategy)
    B, dup_b = _dedupe(B, dup_strategy)

    m = pd.merge(A, B, on="key", how="outer", suffixes=("_a", "_b"), indicator=True)
    m["Match"] = m["_merge"].map({"both": "Matched", "left_only": "Only in A",
                                  "right_only": "Only in B"}).astype(str)

    def _why(r):
        why = []
        for side, lab in (("a", "A"), ("b", "B")):
            rs = r.get(f"reason_{side}")
            if isinstance(rs, str) and rs:
                raw = r.get(f"raw_{side}", "")
                why.append(f"{lab}: {_REASON_TXT.get(rs, rs)} ({raw})" if raw else
                           f"{lab}: {_REASON_TXT.get(rs, rs)}")
        return "; ".join(why)

    an_lbl = analyte_a if analyte_a == analyte_b else f"{analyte_a} ↔ {analyte_b}"
    rep = pd.DataFrame({
        "SampleID": m["id_orig_a"].where(m["id_orig_a"].notna(), m["id_orig_b"]).astype(str).str.strip(),
        "Analyte": an_lbl,
        label_a: m["value_a"].astype(float),
        label_b: m["value_b"].astype(float),
        "Original A": m["raw_a"].fillna(""),
        "Original B": m["raw_b"].fillna(""),
        "Match": m["Match"],
    })
    rep["Note"] = [(_why(r) if r["Match"] == "Matched" else "") for _, r in m.iterrows()]
    flags = []
    for _, r in m.iterrows():
        f = [f"{s.upper()}: {r[f'flag_{s}']}" for s in ("a", "b")
             if isinstance(r.get(f"flag_{s}"), str) and r.get(f"flag_{s}")]
        flags.append("; ".join(f))
    rep["Flags"] = flags
    rep = rep.sort_values(["Match", "SampleID"], kind="stable").reset_index(drop=True)

    matched = rep[rep["Match"] == "Matched"]
    xa = matched[label_a].values.astype(float); yb = matched[label_b].values.astype(float)
    ok = np.isfinite(xa) & np.isfinite(yb)

    excl = Counter()
    for _, r in m[m["Match"] == "Matched"].iterrows():
        for side, lab in (("a", "A"), ("b", "B")):
            rs = r[f"reason_{side}"]
            if isinstance(rs, str) and rs:
                excl[f"{lab}: {_REASON_TXT.get(rs, rs)}"] += 1

    summary = {
        "matched_ids": int(len(matched)), "used": int(ok.sum()),
        "excluded": int((~ok).sum()), "excluded_reasons": dict(excl),
        "only_a": int((rep["Match"] == "Only in A").sum()),
        "only_b": int((rep["Match"] == "Only in B").sum()),
        "dup_removed_a": dup_a, "dup_removed_b": dup_b,
        "sci_ids_a": ia["n_sci"], "sci_ids_b": ib["n_sci"],
        "blank_ids_a": ia["n_blank_id"], "blank_ids_b": ib["n_blank_id"],
        "parse_a": ia["parse"], "parse_b": ib["parse"],
        "analyte_a": analyte_a, "analyte_b": analyte_b,
        "summary_rows_a": ia["n_summary"], "summary_rows_b": ib["n_summary"],
        "summary_ids": sorted(set(ia["summary_ids"]) | set(ib["summary_ids"])),
        "unit_factor": unit_factor(xa[ok], yb[ok]),
    }
    return xa[ok], yb[ok], rep, summary


def _to_datetime(series: pd.Series) -> pd.Series:
    """Tolka datum/tid i vanliga exportformat; NaT där det inte går."""
    s = series.astype(str).str.strip()
    for kw in ({"format": "ISO8601"}, {"dayfirst": False}, {"dayfirst": True}):
        try:
            out = pd.to_datetime(s, errors="coerce", **kw)
            if out.notna().all():
                return out
        except (ValueError, TypeError):
            continue
    return pd.to_datetime(s, errors="coerce")


def extract_precision_replicates(df, id_col, result_col, sample_id,
                                  sort_col=None, analysis_col=None,
                                  analyte=None, n_per_day=5,
                                  group_col=None):
    """
    group_col : if given, rows are grouped by the DISTINCT VALUES of this column
                (e.g. a Date column) -> one day per value. This is the correct
                approach when the file records the real run date, because the
                number of replicates per day may vary between analytes.
                When None, rows are sliced every `n_per_day` in sort order.
    """
    _k, _ = normalize_ids(df[id_col]); _t, _ = normalize_ids(pd.Series([sample_id]))
    mask = (_k == _t.iloc[0]).values
    sub = df[mask].copy()
    if analysis_col and analyte and analyte not in ("ALL", "", None):
        sub = sub[sub[analysis_col].astype(str).str.strip() == str(analyte).strip()]
    if sub.empty:
        raise ValueError(f"No rows found for SampleID = '{sample_id}'.")
    if sort_col and sort_col in sub.columns:
        try:
            _dt = _to_datetime(sub[sort_col])
            if _dt.notna().all():
                sub[sort_col] = _dt
        except Exception:
            pass
        sub = sub.sort_values(sort_col)
    sub = sub.reset_index(drop=True)
    sub[result_col] = _coerce_numeric(sub[result_col])
    sub = sub.dropna(subset=[result_col]).reset_index(drop=True)
    # ── Group by a real day/date column ──────────────────────────────────────
    if group_col is not None and group_col in sub.columns:
        _keys = sub[group_col]
        if not pd.api.types.is_datetime64_any_dtype(_keys):
            _dt = _to_datetime(_keys)                 # '2026-09-15 06:12' -> datum
            if _dt.notna().all():
                _keys = _dt
        if pd.api.types.is_datetime64_any_dtype(_keys):
            _keys = _keys.dt.date
        sub = sub.copy()
        sub["_daykey"] = _keys.astype(str)
        _order = list(dict.fromkeys(sub["_daykey"]))   # preserve sorted order
        data_dict, counts = {}, []
        for i, k in enumerate(_order, start=1):
            vals = sub.loc[sub["_daykey"] == k, result_col].tolist()
            data_dict[f"Day {i}"] = vals
            counts.append(len(vals))
        if len(data_dict) < 2:
            raise ValueError(
                f"Only {len(data_dict)} distinct value(s) in '{group_col}' — "
                "need at least 2 days.")
        if min(counts) < 2:
            raise ValueError("Each day needs at least 2 replicates.")
        n_trim = 0
        sub["Used"] = True
        if min(counts) != max(counts):
            # EP15-beräkningen kräver lika många replikat per dag. Varje dag
            # begränsas till de m första resultaten (i tidsordning); resten
            # markeras och RAPPORTERAS, i stället för att tyst försvinna.
            m = min(counts)
            data_dict = {k: v[:m] for k, v in data_dict.items()}
            for k in _order:
                idx = sub.index[sub["_daykey"] == k]
                sub.loc[idx[m:], "Used"] = False
            n_trim = int(sum(counts) - m * len(counts))
        sub["Day"] = sub["_daykey"].map(
            {k: f"Day {i}" for i, k in enumerate(_order, start=1)})
        sub = sub.drop(columns=["_daykey"])
        return data_dict, sub, n_trim

    # ── Otherwise: slice every n_per_day rows in sort order ──────────────────
    if len(sub) < n_per_day:
        raise ValueError(f"Only {len(sub)} valid results — need at least {n_per_day} (n_per_day).")
    n_complete = len(sub) // n_per_day
    sub["Day"] = [f"Day {i // n_per_day + 1}" if i < n_complete * n_per_day else "Incomplete"
                  for i in range(len(sub))]
    sub_complete = sub[sub["Day"] != "Incomplete"].copy()
    data_dict = {}
    for d in range(n_complete):
        lbl = f"Day {d+1}"
        vals = sub_complete.iloc[d*n_per_day:(d+1)*n_per_day][result_col].tolist()
        data_dict[lbl] = vals
    n_leftover = len(sub) % n_per_day
    return data_dict, sub_complete, n_leftover


def build_matched_excel(report_df, label_a="Method A", label_b="Method B", analyte="") -> bytes:
    import openpyxl
    from openpyxl.styles import PatternFill, Font, Alignment
    from openpyxl.utils.dataframe import dataframe_to_rows
    GREEN     = PatternFill("solid", fgColor="D1FAE5")   # included
    RED       = PatternFill("solid", fgColor="FECACA")   # excluded
    YELLOW    = PatternFill("solid", fgColor="FEF3C7")   # unmatched
    HDR_FILL  = PatternFill("solid", fgColor="1E40AF")
    HDR_FONT  = Font(color="FFFFFF", bold=True)
    RED_FONT  = Font(color="991B1B", bold=True)          # dark red, bold
    GREEN_FONT= Font(color="065F46")                     # dark green

    def _sheet(ws, df, fill, title, status_aware=False):
        """status_aware: colour each row by its Status column
           (Excluded -> red, otherwise green) instead of one flat fill."""
        ws.title = title
        if df.empty:
            ws.append([f"No rows: {title}"]); return

        cols = list(df.columns)
        status_i = cols.index("Status") + 1 if (status_aware and "Status" in cols) else None

        for ri, row in enumerate(dataframe_to_rows(df, index=False, header=True), 1):
            ws.append(row)
            if ri == 1:
                for cell in ws[ri]:
                    cell.fill = HDR_FILL
                    cell.font = HDR_FONT
                    cell.alignment = Alignment(horizontal="center")
                continue

            row_fill, row_font = fill, None
            if status_i is not None:
                val = ws.cell(row=ri, column=status_i).value
                v_ = str(val).strip().lower()
                if v_ == "excluded":
                    row_fill, row_font = RED, RED_FONT
                elif v_.startswith("not used"):
                    row_fill, row_font = YELLOW, None
                else:
                    row_fill, row_font = GREEN, GREEN_FONT

            for cell in ws[ri]:
                cell.fill = row_fill
                if row_font is not None:
                    cell.font = row_font

        ws.freeze_panes = "A2"
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(
                max(len(str(c.value or "")) for c in col) + 4, 40)

    wb = openpyxl.Workbook()
    matched = report_df[report_df["Match"]=="Matched"].drop(columns=["Match"])
    only_a  = report_df[report_df["Match"]=="Only in A"].drop(columns=[label_b,"Match"], errors="ignore")
    only_b  = report_df[report_df["Match"]=="Only in B"].drop(columns=[label_a,"Match"], errors="ignore")
    _sheet(wb.active, matched, GREEN, "Matched pairs", status_aware=True)
    _sheet(wb.create_sheet("Only in A"), only_a,  YELLOW, "Only in A")
    _sheet(wb.create_sheet("Only in B"), only_b,  YELLOW, "Only in B")
    buf = io.BytesIO(); wb.save(buf); buf.seek(0); return buf.read()


def build_precision_excel(raw_df, pr_results, sample_id, analyte, decimals=4,
                          t=None) -> bytes:
    if t is None:
        t = lambda x: x
    import openpyxl
    from openpyxl.styles import PatternFill, Font, Alignment
    from openpyxl.utils.dataframe import dataframe_to_rows
    HDR_FILL = PatternFill("solid", fgColor="1E40AF")
    HDR_FONT = Font(color="FFFFFF", bold=True)
    BLUE     = PatternFill("solid", fgColor="EFF6FF")
    GREEN    = PatternFill("solid", fgColor="ECFDF5")

    def _f(v): return f"{v:.{decimals}f}".replace(".",",")

    wb = openpyxl.Workbook()
    ws1 = wb.active; ws1.title = t("Raw replicates")[:31]
    for ri, row in enumerate(dataframe_to_rows(raw_df, index=False, header=True), 1):
        ws1.append(row)
        for cell in ws1[ri]:
            if ri==1: cell.fill=HDR_FILL; cell.font=HDR_FONT; cell.alignment=Alignment(horizontal="center")
            else: cell.fill=BLUE
    for col in ws1.columns:
        ws1.column_dimensions[col[0].column_letter].width=18

    # ── Sheet 2: sammanfattning i laboratoriets tabellformat ─────────────
    from openpyxl.styles import Border, Side

    ws2 = wb.create_sheet(t("Summary")[:31])

    HDR_BG   = PatternFill("solid", fgColor="000000")   # svart rubrikrad
    LBL_BG   = PatternFill("solid", fgColor="DCE6F1")   # ljusbla etiketter
    VAL_BG   = PatternFill("solid", fgColor="FFFFFF")
    PREC_BG  = PatternFill("solid", fgColor="FFF2CC")   # markera precisionsraderna
    WHITE_B  = Font(color="FFFFFF", bold=True, size=11)
    LBL_F    = Font(bold=True, size=11)
    VAL_F    = Font(size=11)
    thin     = Side(style="thin", color="000000")
    BORDER   = Border(left=thin, right=thin, top=thin, bottom=thin)
    CENTER   = Alignment(horizontal="center", vertical="center")
    RIGHT    = Alignment(horizontal="right",  vertical="center")

    def _pct(v):
        return f"{v:.1f}".replace(".", ",") + "%"

    title_left  = analyte  or t("Analysis")
    title_right = sample_id or t("Control")

    rows = [
        (t("Count:"),                 str(pr_results["n_total_meas"]),      False),
        (t("Mean:"),                  _f(pr_results["grand_mean"]),         False),
        ("SD:",                       _f(pr_results.get("pooled_sd", 0.0)), False),
        ("CV%:",                      _pct(pr_results.get("pooled_cv", 0.0)), False),
        ("Min:",                      _f(pr_results.get("overall_min", 0.0)), False),
        ("Max:",                      _f(pr_results.get("overall_max", 0.0)), False),
        (t("Within-run precision CV%:"), _pct(pr_results["cv_r"]),             True),
        (t("Total imprecision CV%:"),  _pct(pr_results["cv_l"]),             True),
    ]

    # Rubrikrad
    ws2.cell(row=1, column=1, value=title_left)
    ws2.cell(row=1, column=2, value=title_right)
    for c in (1, 2):
        cell = ws2.cell(row=1, column=c)
        cell.fill = HDR_BG; cell.font = WHITE_B
        cell.alignment = CENTER; cell.border = BORDER
    ws2.row_dimensions[1].height = 30

    # Datarader
    for i, (label, value, is_prec) in enumerate(rows, start=2):
        lc = ws2.cell(row=i, column=1, value=label)
        vc = ws2.cell(row=i, column=2, value=value)
        lc.fill = PREC_BG if is_prec else LBL_BG
        vc.fill = PREC_BG if is_prec else VAL_BG
        lc.font = LBL_F;  vc.font = LBL_F if is_prec else VAL_F
        lc.alignment = RIGHT; vc.alignment = CENTER
        lc.border = BORDER;   vc.border = BORDER
        ws2.row_dimensions[i].height = 19

    ws2.column_dimensions["A"].width = 26
    ws2.column_dimensions["B"].width = 18

    # Fotnot med programversion och uppläggning
    fr = len(rows) + 3
    try:
        import sys as _sys, os as _os
        _sys.path.insert(0, _os.path.dirname(_os.path.dirname(_os.path.abspath(__file__))))
        from version import stamp as _stamp
        _v = _stamp()
    except Exception:
        _v = ""
    notes = [
        t("Design: {D} days x {n} replicates ({N} measurements)")
        .format(D=pr_results['D'], n=pr_results['n'], N=pr_results['n_total_meas']),
        t("SD and CV% refer to all measurements pooled."),
        t("Within-run precision and total imprecision according to CLSI EP15-A3."),
        _v,
    ]
    for k, txt in enumerate(notes):
        if not txt:
            continue
        c = ws2.cell(row=fr + k, column=1, value=txt)
        c.font = Font(size=8, italic=True, color="595959")

    buf = io.BytesIO(); wb.save(buf); buf.seek(0); return buf.read()
